"""md / txt / docx documentation uploads."""
import os

import docx

from app.adapters.base import FAILED, OK, AdapterResult, register
from app.adapters.common import MAX_FILE_BYTES


@register("documentation")
def fetch_documentation(artifact, storage) -> AdapterResult:
    if not artifact.storage_key:
        return AdapterResult(FAILED, "no uploaded file")
    name = os.path.basename(artifact.source_ref or "doc.txt")
    ext = os.path.splitext(name)[1].lower()
    path = storage.local_path(artifact.storage_key)
    if ext in (".md", ".txt", ".rst", ""):
        text = open(path, "rb").read(MAX_FILE_BYTES).decode("utf-8", errors="replace")
    elif ext == ".docx":
        try:
            d = docx.Document(path)
        except Exception as e:  # noqa: BLE001
            return AdapterResult(FAILED, f"unreadable docx: {type(e).__name__}")
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        text = "\n".join(parts)
    else:
        return AdapterResult(FAILED, f"unsupported documentation type {ext!r} (use md, txt or docx)")
    if not text.strip():
        return AdapterResult(FAILED, "document is empty")
    return AdapterResult(OK, "", {name: text})
